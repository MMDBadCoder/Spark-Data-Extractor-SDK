import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path(__file__).resolve().parent
text=(root/'docs/customer_guide_fa.md').read_text()
doc=Document(); section=doc.sections[0]
section.page_width=Inches(8.5); section.page_height=Inches(11)
section.top_margin=section.bottom_margin=Inches(.8)
section.left_margin=section.right_margin=Inches(.8)
for name in ['Normal','Title','Heading 1','Heading 2']:
 style=doc.styles[name]; style.font.name='DejaVu Sans'; style.font.color.rgb=RGBColor(0,0,0)
 style.font.size=Pt(11 if name=='Normal' else 19 if name=='Title' else 14)
 style.paragraph_format.space_after=Pt(8)
 style.paragraph_format.line_spacing=1.18
code=False
for line in text.splitlines():
 if line.startswith('```'): code=not code; continue
 if not line: continue
 style='Normal' if code else 'Title' if line.startswith('# ') else 'Heading 1' if line.startswith('## ') else 'Normal'
 p=doc.add_paragraph(line.lstrip('# ') if style!='Normal' else line,style)
 p.alignment=WD_ALIGN_PARAGRAPH.LEFT if code else WD_ALIGN_PARAGRAPH.RIGHT
 if not code:
  bidi=OxmlElement('w:bidi'); bidi.set(qn('w:val'),'1'); p._p.get_or_add_pPr().append(bidi)
 for r in p.runs:
  r.font.name='DejaVu Sans Mono' if code else 'DejaVu Sans'
  r.font.size=Pt(9 if code else 19 if style=='Title' else 14 if style=='Heading 1' else 11)
  if not code:
   rtl=OxmlElement('w:rtl'); rtl.set(qn('w:val'),'1'); r._r.get_or_add_rPr().append(rtl)
  fonts=r._r.get_or_add_rPr().get_or_add_rFonts(); fonts.set(qn('w:cs'),'DejaVu Sans')
 p.paragraph_format.keep_together=True
 if code: p.paragraph_format.space_after=Pt(2)
for style in doc.styles:
 for border in list(style.element.iter(qn('w:pBdr'))):
  border.getparent().remove(border)
for p in doc.paragraphs:
 for border in list(p._p.iter(qn('w:pBdr'))):
  border.getparent().remove(border)
doc.core_properties.title='راهنمای استخراج داده در دفترچه مشتری'
doc.core_properties.author='Data Extractor'
doc.save(root/'docs/customer_guide_fa.docx')
cells=[]
def md(s): cells.append({'cell_type':'markdown','metadata':{},'source':s.splitlines(True)})
def codecell(s): cells.append({'cell_type':'code','execution_count':None,'metadata':{},'outputs':[],'source':s.splitlines(True)})
md('# استخراج داده\nمدیر اتصال را آماده کرده است. نام ستون‌های نمونه را با ستون‌های واقعی جایگزین کنید. تاریخ‌ها حداکثر هفت روز و شامل ابتدا و پایان هستند. هر بار یک کار اجرا کنید. راهنما: docs/customer_guide_fa.docx')
codecell('from pyspark.sql.functions import col\nfrom data_extractor import Query\n\nquery = Query("customer_search")\nquery.between("2026-01-01", "2026-01-07")\nquery.where((col("a") > 10) & col("b").startswith("asdasa"))\nquery.select("a", "b", "c")\nquery.limit(1_000_000)\nquery.output_files(4)\nquery.run()\n')
md('## شرط‌ها\n& یعنی و، | یعنی یا و ~ یعنی نقیض. هر مقایسه را در پرانتز بنویسید؛ and و or پایتون مناسب نیستند. چند where با و ترکیب می‌شوند. تعداد فایل هدف تضمین تعداد واقعی فایل نیست. انتخاب رکوردهای limit ترتیب تضمین‌شده ندارد.')
md('## دریافت\nفرمان دریافت چاپ‌شده در پیام پایان را روی دستگاه لینوکسی آماده اجرا کنید. تاریخ فرمان تاریخ اجرای تهران است، نه تاریخ ورودی. پس از خطای مهلت، کار را مکرراً تکرار نکنید؛ قفل ممکن است تا توقف کار نگه داشته شود. برای گزارش خطا شناسه و زمان را به مدیر بدهید.')
(root/'notebooks/customer_example.ipynb').write_text(json.dumps({'cells':cells,'metadata':{'kernelspec':{'display_name':'Python 3','language':'python','name':'python3'},'language_info':{'name':'python','version':'3.10'}},'nbformat':4,'nbformat_minor':5},ensure_ascii=False,indent=2))
